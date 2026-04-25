from __future__ import annotations

import argparse
import re
import shutil
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"

ET.register_namespace("w", W_NS)
ET.register_namespace("r", R_NS)
ET.register_namespace("", PR_NS)
ET.register_namespace("", CT_NS)


def get_markdown_body_text(content: str) -> str:
    normalized = content.replace("\ufeff", "").replace("\r", "")
    if normalized.startswith("---\n"):
        match = re.match(r"(?s)^---\n.*?\n---\n?", normalized)
        if match:
            return normalized[match.end() :].strip()

    lines = normalized.split("\n")
    metadata_pattern = re.compile(r"^[A-Za-z][A-Za-z0-9_]*:\s*.+$")
    metadata_count = 0
    body_start = 0

    for index, line in enumerate(lines):
        stripped = line.strip()
        if not stripped:
            if metadata_count >= 3:
                body_start = index + 1
                break
            metadata_count = 0
            break
        if metadata_pattern.match(stripped):
            metadata_count += 1
            continue
        metadata_count = 0
        break

    if metadata_count >= 3 and body_start > 0:
        return "\n".join(lines[body_start:]).strip()

    return normalized.strip()


def get_front_matter_value(content: str, key: str) -> str | None:
    normalized = content.replace("\ufeff", "").replace("\r", "")
    if not normalized.startswith("---\n"):
        return None
    match = re.match(r"(?s)^---\n(.*?)\n---\n?", normalized)
    if not match:
        return None
    value_match = re.search(rf"(?m)^{re.escape(key)}:\s*(.+?)\s*$", match.group(1))
    if not value_match:
        return None
    return value_match.group(1).strip().strip('"')


def convert_markdown_line_to_plain_text(line: str) -> str:
    text = line.replace("\t", "    ")
    text = re.sub(r"^\s*>\s?", "", text)
    text = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("```", "")
    text = text.replace("`", "")
    text = text.replace("***", "")
    text = text.replace("**", "")
    text = text.replace("*", "")
    text = text.replace("___", "")
    text = text.replace("__", "")
    text = text.replace("_", "")
    text = text.replace("~~", "")
    return text.rstrip()


def get_markdown_paragraph_specs(content: str) -> list[tuple[str, str]]:
    body = get_markdown_body_text(content)
    lines = body.replace("\r", "").split("\n")
    paragraphs: list[tuple[str, str]] = []
    in_code_block = False

    for original_line in lines:
        line = original_line
        trimmed = line.strip()

        if trimmed.startswith("```") or trimmed.startswith("~~~"):
            in_code_block = not in_code_block
            continue

        if in_code_block:
            if line.strip():
                paragraphs.append(("Normal", line.rstrip()))
            continue

        if re.match(r"^\s*---+\s*$", trimmed) or re.match(r"^\s*\*\*\*+\s*$", trimmed):
            continue

        if not trimmed:
            paragraphs.append(("Blank", ""))
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", trimmed)
        if heading_match:
            heading_level = min(len(heading_match.group(1)), 3)
            text = convert_markdown_line_to_plain_text(heading_match.group(2)).strip()
            paragraphs.append(("Normal", text))
            continue

        bullet_match = re.match(r"^\s*[-*+]\s+(.+)$", trimmed)
        if bullet_match:
            text = convert_markdown_line_to_plain_text(bullet_match.group(1)).strip()
            paragraphs.append(("Normal", f"• {text}"))
            continue

        if re.match(r"^\s*\d+\.\s+(.+)$", trimmed):
            text = convert_markdown_line_to_plain_text(trimmed).strip()
            paragraphs.append(("Normal", text))
            continue

        paragraphs.append(("Normal", convert_markdown_line_to_plain_text(line).strip()))

    return paragraphs


def extract_chapter_title(content: str, fallback: str) -> str:
    front_matter_title = get_front_matter_value(content, "title")
    if front_matter_title:
        return front_matter_title.strip()

    body = get_markdown_body_text(content)
    for line in body.replace("\r", "").split("\n"):
        trimmed = line.strip()
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", trimmed)
        if heading_match:
            title = convert_markdown_line_to_plain_text(heading_match.group(2)).strip()
            if title:
                return title

    return fallback


def sort_key(path: Path) -> tuple[int, str]:
    match = re.match(r"^(\d+)", path.stem)
    return (int(match.group(1)) if match else 9999, path.name.lower())


def build_simple_docx(source_root: Path, output_file: Path, title: str, author: str, include_toc: bool = False) -> None:
    chapter_root = source_root / "02_chapters"
    md_files = sorted(
        [path for path in chapter_root.glob("*.md") if not path.name.startswith("_")],
        key=sort_key,
    )
    if not md_files:
        raise RuntimeError(f"No markdown files found under {chapter_root}")

    document = Document()
    core = document.core_properties
    core.title = title or ""
    core.author = author or ""
    core.last_modified_by = author or ""
    core.subject = ""
    core.category = ""
    core.comments = ""
    core.keywords = ""

    for section in document.sections:
        section.header.paragraphs[0].text = ""
        section.footer.paragraphs[0].text = ""

    chapter_contents: list[tuple[Path, str, str]] = []
    for md_file in md_files:
        content = md_file.read_text(encoding="utf-8")
        chapter_title = extract_chapter_title(content, md_file.stem)
        chapter_contents.append((md_file, content, chapter_title))

    if include_toc:
        toc_heading = "目录" if re.match(r"^(zh|zh-tw)$", source_root.name, re.IGNORECASE) else "Contents"
        toc_title = document.add_paragraph()
        toc_title.style = "Normal"
        toc_title.add_run(toc_heading)
        document.add_paragraph()
        for index, (_, _, chapter_title) in enumerate(chapter_contents, start=1):
            toc_entry = document.add_paragraph()
            toc_entry.style = "Normal"
            toc_entry.add_run(f"{index}. {chapter_title}")
        document.add_page_break()

    for md_file, content, _chapter_title in chapter_contents:
        for style, text in get_markdown_paragraph_specs(content):
            paragraph = document.add_paragraph()
            if style != "Blank":
                try:
                    paragraph.style = style
                except Exception:
                    paragraph.style = "Normal"
            if text:
                paragraph.add_run(text)
        document.add_paragraph()

    output_file.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_file))
    sanitize_docx_package(output_file)


def sanitize_docx_package(output_file: Path) -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_root = Path(temp_dir)
        with zipfile.ZipFile(output_file, "r") as archive:
            archive.extractall(temp_root)

        document_path = temp_root / "word" / "document.xml"
        document_rels_path = temp_root / "word" / "_rels" / "document.xml.rels"
        content_types_path = temp_root / "[Content_Types].xml"
        package_rels_path = temp_root / "_rels" / ".rels"
        custom_xml_root = temp_root / "customXml"
        header_path = temp_root / "word" / "header1.xml"
        footer_path = temp_root / "word" / "footer1.xml"

        if document_path.exists():
            tree = ET.parse(document_path)
            root = tree.getroot()
            for sect in root.findall(f".//{{{W_NS}}}sectPr"):
                for tag in ("headerReference", "footerReference", "titlePg", "pgNumType"):
                    for child in list(sect.findall(f"{{{W_NS}}}{tag}")):
                        sect.remove(child)
            tree.write(document_path, encoding="utf-8", xml_declaration=True)

        if document_rels_path.exists():
            tree = ET.parse(document_rels_path)
            root = tree.getroot()
            for rel in list(root):
                rel_type = rel.attrib.get("Type", "")
                target = rel.attrib.get("Target", "")
                if (
                    rel_type.endswith("/header")
                    or rel_type.endswith("/footer")
                    or rel_type.endswith("/customXml")
                    or rel_type.endswith("/numbering")
                    or rel_type.endswith("/webSettings")
                    or rel_type.endswith("/stylesWithEffects")
                    or target in {"header1.xml", "footer1.xml", "../customXml/item1.xml", "numbering.xml", "webSettings.xml", "stylesWithEffects.xml"}
                ):
                    root.remove(rel)
            tree.write(document_rels_path, encoding="utf-8", xml_declaration=True)

        if content_types_path.exists():
            tree = ET.parse(content_types_path)
            root = tree.getroot()
            for node in list(root):
                part_name = node.attrib.get("PartName", "")
                if part_name in {
                    "/word/header1.xml",
                    "/word/footer1.xml",
                    "/customXml/item1.xml",
                    "/customXml/itemProps1.xml",
                    "/docProps/thumbnail.jpeg",
                    "/word/numbering.xml",
                    "/word/webSettings.xml",
                    "/word/stylesWithEffects.xml",
                }:
                    root.remove(node)
            tree.write(content_types_path, encoding="utf-8", xml_declaration=True)

        if package_rels_path.exists():
            tree = ET.parse(package_rels_path)
            root = tree.getroot()
            for rel in list(root):
                rel_type = rel.attrib.get("Type", "")
                target = rel.attrib.get("Target", "")
                if rel_type.endswith("/customXml") or target.startswith("customXml/"):
                    root.remove(rel)
            tree.write(package_rels_path, encoding="utf-8", xml_declaration=True)

        if header_path.exists():
            header_path.unlink()
        if footer_path.exists():
            footer_path.unlink()
        if custom_xml_root.exists():
            shutil.rmtree(custom_xml_root, ignore_errors=True)

        for extra in [
            temp_root / "word" / "numbering.xml",
            temp_root / "word" / "webSettings.xml",
            temp_root / "word" / "stylesWithEffects.xml",
        ]:
            if extra.exists():
                extra.unlink()

        rel_item = temp_root / "customXml" / "_rels" / "item1.xml.rels"
        if rel_item.exists():
            rel_item.unlink()

        thumb = temp_root / "docProps" / "thumbnail.jpeg"
        if thumb.exists():
            thumb.unlink()

        with zipfile.ZipFile(output_file, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in sorted(temp_root.rglob("*")):
                if path.is_file():
                    archive.write(path, path.relative_to(temp_root).as_posix())


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", required=True)
    parser.add_argument("--output-file", required=True)
    parser.add_argument("--title", default="")
    parser.add_argument("--author", default="")
    parser.add_argument("--include-toc", action="store_true")
    args = parser.parse_args()

    build_simple_docx(
        source_root=Path(args.source_root),
        output_file=Path(args.output_file),
        title=args.title,
        author=args.author,
        include_toc=args.include_toc,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
